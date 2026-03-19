"""
Excel Generator - Takes parsed PDF data and creates a consolidated Excel report.
"""

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

logger = logging.getLogger(__name__)

# Style constants
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SUBHEADER_FONT = Font(bold=True, size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


class ExcelGenerator:
    """Generates a consolidated Excel report from parsed PDF data."""

    def generate(self, parsed_data: list[dict], output_path: Path = None) -> Path:
        """Generate an Excel file from parsed PDF data.

        Args:
            parsed_data: list of parsed PDF results from PDFParser
            output_path: optional output file path

        Returns:
            Path to the generated Excel file
        """
        if output_path is None:
            today = datetime.now().strftime("%Y-%m-%d")
            filename = config.EXCEL_FILENAME_TEMPLATE.format(date=today)
            output_path = config.OUTPUT_DIR / filename

        output_path.parent.mkdir(parents=True, exist_ok=True)

        logger.info("Generating Excel report: %s", output_path)

        wb = Workbook()
        # Remove the default sheet
        wb.remove(wb.active)

        # 1. Create Summary sheet
        self._create_summary_sheet(wb, parsed_data)

        # 2. Create a single consolidated sheet with all report table data
        self._create_consolidated_sheet(wb, parsed_data)

        # 3. Create a raw text sheet for reference
        self._create_text_sheet(wb, parsed_data)

        wb.save(str(output_path))
        logger.info("Excel report saved: %s", output_path)
        return output_path

    def _create_summary_sheet(self, wb: Workbook, parsed_data: list[dict]):
        """Create a summary sheet with an overview of all downloaded reports."""
        ws = wb.create_sheet("Summary")

        # Title row
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"ICE Market Reports - {datetime.now().strftime('%Y-%m-%d')}"
        title_cell.font = Font(bold=True, size=14, color="1F4E79")
        title_cell.alignment = Alignment(horizontal="center")

        # Headers
        headers = [
            "Report ID", "Report Name", "Filename",
            "Pages", "Tables Found", "Download Time", "File Size (KB)"
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row_idx, report in enumerate(parsed_data, 4):
            dl_info = report.get("download_info", {})
            meta = report.get("metadata", {})
            values = [
                dl_info.get("report_id", ""),
                dl_info.get("report_name", meta.get("title", "")),
                report.get("filename", ""),
                meta.get("page_count", 0),
                len(report.get("tables", [])),
                dl_info.get("download_time", ""),
                round(dl_info.get("file_size", 0) / 1024, 1),
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(horizontal="center")

        # Auto-fit column widths
        self._auto_fit_columns(ws)

    def _create_consolidated_sheet(self, wb: Workbook, parsed_data: list[dict]):
        """Create a single sheet with all report table data consolidated."""
        ws = wb.create_sheet("All Reports Data")

        # Title
        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = f"ICE Market Reports Data - {datetime.now().strftime('%Y-%m-%d')}"
        title_cell.font = Font(bold=True, size=14, color="1F4E79")
        title_cell.alignment = Alignment(horizontal="center")

        current_row = 3

        for idx, report in enumerate(parsed_data):
            tables = report.get("tables", [])
            if not tables:
                continue

            dl_info = report.get("download_info", {})
            report_name = dl_info.get("report_name", report.get("filename", f"Report_{idx}"))

            # Report section header
            ws.merge_cells(f"A{current_row}:H{current_row}")
            cell = ws.cell(row=current_row, column=1, value=report_name)
            cell.font = Font(bold=True, size=12, color="FFFFFF")
            cell.fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
            cell.alignment = Alignment(horizontal="left")
            current_row += 1

            for table in tables:
                # Column headers
                headers = table.get("headers", [])
                if headers:
                    for col, h in enumerate(headers, 1):
                        cell = ws.cell(row=current_row, column=col, value=h)
                        cell.font = HEADER_FONT
                        cell.fill = HEADER_FILL
                        cell.border = THIN_BORDER
                    current_row += 1

                # Data rows
                for row_data in table.get("rows", []):
                    for col, val in enumerate(row_data, 1):
                        cell = ws.cell(row=current_row, column=col, value=self._try_numeric(val))
                        cell.border = THIN_BORDER
                    current_row += 1

                # Small gap between tables from the same report
                current_row += 1

            # Larger gap between reports
            current_row += 1

        self._auto_fit_columns(ws)

    def _create_text_sheet(self, wb: Workbook, parsed_data: list[dict]):
        """Create a sheet with raw text content from all reports."""
        ws = wb.create_sheet("Raw Text")

        ws.merge_cells("A1:C1")
        ws["A1"].value = "Raw Text Content from All Reports"
        ws["A1"].font = Font(bold=True, size=12, color="1F4E79")

        headers = ["Report", "Page", "Text Content"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER

        row = 4
        for report in parsed_data:
            dl_info = report.get("download_info", {})
            name = dl_info.get("report_name", report.get("filename", ""))
            for text_block in report.get("text_content", []):
                ws.cell(row=row, column=1, value=name).border = THIN_BORDER
                ws.cell(row=row, column=2, value=text_block["page"]).border = THIN_BORDER
                # Truncate text to Excel cell limit (32767 chars)
                text = text_block["text"][:32000]
                cell = ws.cell(row=row, column=3, value=text)
                cell.border = THIN_BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                row += 1

        # Set column widths
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 8
        ws.column_dimensions["C"].width = 100

    def _auto_fit_columns(self, ws):
        """Auto-fit column widths based on content."""
        for col_cells in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 4, 50)
            ws.column_dimensions[col_letter].width = adjusted_width

    @staticmethod
    def _try_numeric(value: str):
        """Try to convert a string value to a number for Excel."""
        if not value or not isinstance(value, str):
            return value
        # Remove common formatting
        cleaned = value.replace(",", "").replace("%", "").strip()
        try:
            if "." in cleaned:
                return float(cleaned)
            return int(cleaned)
        except (ValueError, TypeError):
            return value
